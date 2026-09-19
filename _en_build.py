# -*- coding: utf-8 -*-
"""Build the English edition from the Chinese document.

Images, tables and the 25 display equations are carried over untouched; only the
prose is replaced. Paragraphs that contain inline equations are rebuilt segment
by segment so that every oMath element keeps its place (or moves to where English
word order needs it).
"""
import io
import shutil
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from _en_trans1 import SIMPLE as S1, SEG as G1
from _en_trans2 import SIMPLE as S2

SRC = 'PA-STL-DualTimesNet_论文.docx'
DST = 'PA-STL-DualTimesNet_EN.docx'

SIMPLE = dict(S1); SIMPLE.update(S2)
SEG = dict(G1)

M = '{http://schemas.openxmlformats.org/officeDocument/2006/math}'
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'

shutil.copyfile(SRC, DST)
d = Document(DST)
ps = d.paragraphs


def add_run(par, text):
    r = OxmlElement('w:r')
    t = OxmlElement('w:t')
    t.text = text
    t.set(qn('xml:space'), 'preserve')
    r.append(t)
    par._element.append(r)


def clear_content(par):
    """Drop every run / oMath, keeping the paragraph properties."""
    keep = par._element.find(qn('w:pPr'))
    for el in list(par._element):
        if el is keep:
            continue
        par._element.remove(el)


simple_n = seg_n = 0
for i, p in enumerate(ps):
    if i in SIMPLE:
        if not p.runs:
            add_run(p, SIMPLE[i])
        else:
            p.runs[0].text = SIMPLE[i]
            for r in p.runs[1:]:
                r.text = ''
        simple_n += 1
    elif i in SEG:
        segs, order = SEG[i]
        maths = list(p._element.findall(M + 'oMath'))
        assert len(order) == len(maths), 'P%d: %d maths, order expects %d' % (i, len(maths), len(order))
        assert len(segs) == len(order) + 1, 'P%d: %d segments for %d maths' % (i, len(segs), len(order))
        clear_content(p)
        for k, s in enumerate(segs):
            if s:
                add_run(p, s)
            if k < len(order):
                p._element.append(maths[order[k]])
        seg_n += 1

print('translated: %d plain + %d with equations' % (simple_n, seg_n))

# ---- captions, to match the reference document ------------------------------
cap_fig = cap_tab = 0
for p in d.paragraphs:
    t = p.text.strip()
    if re.match(r'^Fig\.\s*\d+\s', t) and len(t) < 90 and '。' not in t:
        new = re.sub(r'^Fig\.\s*(\d+)\s*', r'Figure \1. ', t)
        p.runs[0].text = new
        for r in p.runs[1:]:
            r.text = ''
        cap_fig += 1
    elif re.match(r'^Table\s*\d+\s', t) and len(t) < 90 and '。' not in t:
        new = re.sub(r'^Table\s*(\d+)\s*', r'Table \1  ', t)
        p.runs[0].text = new
        for r in p.runs[1:]:
            r.text = ''
        cap_tab += 1
print('captions: %d figure, %d table' % (cap_fig, cap_tab))

d.save(DST)

# ---- report -----------------------------------------------------------------
chk = Document(DST)
cjk = re.compile(r'[一-鿿]')
left = [(i, p.text.strip()[:50]) for i, p in enumerate(chk.paragraphs) if cjk.search(p.text)]
out = io.open('_en_report.txt', 'w', encoding='utf-8')
out.write('remaining Chinese paragraphs: %d\n' % len(left))
for i, t in left:
    out.write('  P%d %s\n' % (i, t))
out.close()
print('saved', DST)
