# -*- coding: utf-8 -*-
"""Fix two equations in both the Chinese and the English Word files.

(5)  the upright "top" operator was stored as the symbol U+22A4 (⊤); write it as
     the word "Top" so it reads arg Top_k rather than arg ⊤_k.
(24) MASE: the denominator now scales by the seasonal naive forecast, so the
     lag-1 difference becomes a lag-96 one.
"""
import shutil

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

M = '{http://schemas.openxmlformats.org/officeDocument/2006/math}'

FILES = ['PA-STL-DualTimesNet_论文.docx', 'PA-STL-DualTimesNet_EN.docx']


def fix_top(doc):
    n = 0
    for p in doc.paragraphs:
        for om in p._element.iter(M + 'oMath'):
            for t in om.iter(M + 't'):
                if t.text and '⊤' in t.text:
                    t.text = t.text.replace('⊤', 'Top')
                    n += 1
                    # an operator name must be upright, not italic
                    r = t.getparent()
                    rpr = r.find(M + 'rPr')
                    if rpr is None:
                        rpr = OxmlElement('m:rPr')
                        r.insert(0, rpr)
                    if rpr.find(M + 'sty') is None:
                        sty = OxmlElement('m:sty')
                        sty.set(M + 'val', 'p')
                        rpr.insert(0, sty)
    return n


def fix_mase(doc):
    """Only the second sum inside equation (24) carries these three texts."""
    hits = []
    for p in doc.paragraphs:
        for om in p._element.iter(M + 'oMath'):
            texts = [t.text for t in om.iter(M + 't') if t.text]
            joined = ''.join(texts)
            if 'j=2' not in joined or 'j-1' not in joined or 'N-1' not in joined:
                continue
            for t in om.iter(M + 't'):
                if t.text == 'N-1':
                    t.text = 'N-96'; hits.append('N-1->N-96')
                elif t.text == 'j=2':
                    t.text = 'j=97'; hits.append('j=2->j=97')
                elif t.text == 'j-1':
                    t.text = 'j-96'; hits.append('j-1->j-96')
    return hits


for f in FILES:
    shutil.copyfile(f, f.replace('.docx', '.pre_eqfix.docx'))
    d = Document(f)
    a = fix_top(d)
    b = fix_mase(d)
    d.save(f)
    print('%-32s top->Top: %d   MASE: %s' % (f, a, b))
